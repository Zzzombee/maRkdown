#!/usr/bin/env python3
"""Check the minimal local environment and summarize a PPTX before extraction."""

from __future__ import annotations

import argparse
import importlib.metadata
import json
import os
import platform
import shutil
import sys
import winreg
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from PIL import features
from pptx import Presentation


def powerpoint_registered() -> bool:
    if platform.system() != "Windows":
        return False
    try:
        with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, r"PowerPoint.Application\CLSID"):
            return True
    except OSError:
        return False


def inspect_pptx(path: Path) -> dict:
    prs = Presentation(path)
    shape_counts: Counter[str] = Counter()
    notes_slides = 0
    tables = 0
    charts = 0
    pictures = 0
    groups = 0

    def walk(shapes):
        nonlocal tables, charts, pictures, groups
        for shape in shapes:
            shape_counts[str(shape.shape_type)] += 1
            if getattr(shape, "has_table", False):
                tables += 1
            if getattr(shape, "has_chart", False):
                charts += 1
            if str(shape.shape_type).startswith("PICTURE"):
                pictures += 1
            if str(shape.shape_type).startswith("GROUP"):
                groups += 1
                yield from walk(shape.shapes)
            yield shape

    for slide in prs.slides:
        list(walk(slide.shapes))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            if slide.notes_slide.notes_text_frame.text.strip():
                notes_slides += 1

    with ZipFile(path) as archive:
        media = [name for name in archive.namelist() if name.lower().startswith("ppt/media/")]
        extensions = Counter(Path(name).suffix.lower() or "<none>" for name in media)

    return {
        "path": str(path.resolve()),
        "size_bytes": path.stat().st_size,
        "modified_time_ns": path.stat().st_mtime_ns,
        "slides": len(prs.slides),
        "notes_slides_with_text": notes_slides,
        "pictures": pictures,
        "groups": groups,
        "tables": tables,
        "charts": charts,
        "media_extensions": dict(sorted(extensions.items())),
        "shape_types": dict(sorted(shape_counts.items())),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    if not args.pptx.is_file():
        parser.error(f"PPTX not found: {args.pptx}")

    report = {
        "environment": {
            "python": sys.version.split()[0],
            "python_executable": sys.executable,
            "python_pptx": importlib.metadata.version("python-pptx"),
            "pillow": importlib.metadata.version("Pillow"),
            "pillow_libtiff": bool(features.check("libtiff")),
            "powershell": shutil.which("pwsh") or shutil.which("powershell"),
            "powerpoint_registered": powerpoint_registered(),
            "platform": platform.platform(),
        },
        "presentation": inspect_pptx(args.pptx),
    }

    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text + os.linesep, encoding="utf-8")
    print(text)
    return 0 if report["environment"]["powerpoint_registered"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
