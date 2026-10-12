#!/usr/bin/env python3
"""Extract research-relevant PPTX structure into stable JSON."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def text_frame_data(text_frame) -> dict[str, Any]:
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        paragraphs.append(
            {
                "text": paragraph.text,
                "level": paragraph.level,
                "runs": [
                    {
                        "text": run.text,
                        "bold": run.font.bold,
                        "italic": run.font.italic,
                        "underline": run.font.underline,
                    }
                    for run in paragraph.runs
                ],
            }
        )
    return {"text": text_frame.text, "paragraphs": paragraphs}


def table_data(table) -> dict[str, Any]:
    """Preserve merged-cell geometry instead of flattening duplicated values."""
    rows = []
    for row_index, row in enumerate(table.rows):
        cells = []
        for column_index, cell in enumerate(row.cells):
            cells.append(
                {
                    "row": row_index,
                    "column": column_index,
                    "text": cell.text,
                    "is_merge_origin": bool(cell.is_merge_origin),
                    "is_spanned": bool(cell.is_spanned),
                    "row_span": int(cell.span_height),
                    "column_span": int(cell.span_width),
                }
            )
        rows.append(cells)
    return {"row_count": len(table.rows), "column_count": len(table.columns), "rows": rows}


def safe_text_frame(value) -> str | None:
    try:
        return value.text_frame.text
    except Exception:
        return None


def chart_data(chart) -> dict[str, Any]:
    """Extract native chart categories and values without rebuilding the deck."""
    result: dict[str, Any] = {
        "chart_type": str(chart.chart_type),
        "title": safe_text_frame(chart.chart_title) if chart.has_title else None,
        "plots": [],
    }
    for plot_index, plot in enumerate(chart.plots):
        plot_result: dict[str, Any] = {
            "index": plot_index,
            "plot_type": type(plot).__name__,
            "categories": [],
            "series": [],
        }
        try:
            plot_result["categories"] = [category.label for category in plot.categories]
        except Exception as exc:
            plot_result["categories_error"] = f"{type(exc).__name__}: {exc}"

        for series_index, series in enumerate(plot.series):
            series_result: dict[str, Any] = {"index": series_index, "name": series.name}
            try:
                series_result["values"] = list(series.values)
            except Exception as exc:
                series_result["values_error"] = f"{type(exc).__name__}: {exc}"
            try:
                series_result["x_values"] = list(series.x_values)
            except Exception:
                pass
            try:
                series_result["bubble_sizes"] = list(series.bubble_sizes)
            except Exception:
                pass
            plot_result["series"].append(series_result)
        result["plots"].append(plot_result)
    return result


def geometry(shape) -> dict[str, float | int | None]:
    def value(name):
        raw = getattr(shape, name, None)
        return int(raw) if raw is not None else None

    return {
        "left_emu": value("left"),
        "top_emu": value("top"),
        "width_emu": value("width"),
        "height_emu": value("height"),
        "rotation_degrees": getattr(shape, "rotation", None),
    }


def image_data(shape, warnings: list[dict[str, Any]]) -> dict[str, Any] | None:
    try:
        image = shape.image
        blob = image.blob
        return {
            "filename": image.filename,
            "extension": image.ext,
            "content_type": image.content_type,
            "size_pixels": list(image.size),
            "dpi": list(image.dpi),
            "sha256": sha256_bytes(blob),
            "bytes": len(blob),
        }
    except Exception as exc:
        warnings.append(
            {
                "severity": "warning",
                "code": "image-unavailable",
                "shape": shape.name,
                "detail": f"{type(exc).__name__}: {exc}",
            }
        )
        return None


def shape_data(shape, warnings: list[dict[str, Any]], path: list[str]) -> dict[str, Any]:
    result: dict[str, Any] = {
        "name": shape.name,
        "path": path + [shape.name],
        "shape_id": shape.shape_id,
        "shape_type": str(shape.shape_type),
        "geometry": geometry(shape),
    }
    if getattr(shape, "has_text_frame", False):
        result["text"] = text_frame_data(shape.text_frame)
    if getattr(shape, "has_table", False):
        result["table"] = table_data(shape.table)
    if getattr(shape, "has_chart", False):
        try:
            result["chart"] = chart_data(shape.chart)
        except Exception as exc:
            warnings.append(
                {
                    "severity": "error",
                    "code": "chart-extraction-failed",
                    "shape": shape.name,
                    "detail": f"{type(exc).__name__}: {exc}",
                }
            )
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        result["image"] = image_data(shape, warnings)
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        result["children"] = [shape_data(child, warnings, result["path"]) for child in shape.shapes]
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if not args.pptx.is_file():
        parser.error(f"PPTX not found: {args.pptx}")

    prs = Presentation(args.pptx)
    warnings: list[dict[str, Any]] = []
    slides = []
    for slide_number, slide in enumerate(prs.slides, 1):
        notes = None
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text
        slides.append(
            {
                "slide_number": slide_number,
                "notes": notes,
                "shapes": [shape_data(shape, warnings, []) for shape in slide.shapes],
            }
        )

    payload = {
        "schema_version": 1,
        "source": {
            "path": str(args.pptx.resolve()),
            "size_bytes": args.pptx.stat().st_size,
            "modified_time_ns": args.pptx.stat().st_mtime_ns,
            "sha256": sha256_bytes(args.pptx.read_bytes()),
        },
        "slide_size_emu": {"width": int(prs.slide_width), "height": int(prs.slide_height)},
        "slide_count": len(prs.slides),
        "slides": slides,
        "warnings": warnings,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + os.linesep, encoding="utf-8")
    print(json.dumps({"output": str(args.output), "slides": len(slides), "warnings": len(warnings)}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
