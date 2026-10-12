#!/usr/bin/env python3
"""Preserve embedded TIFF evidence and create explicitly non-quantitative PNG previews."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import os
from pathlib import Path
from typing import Any

from PIL import Image, ImageCms
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def image_part(shape):
    relationship_id = shape._pic.blipFill.blip.rEmbed
    if not relationship_id:
        return None
    return shape.part.related_part(relationship_id)


def walk(shapes, path: list[str]):
    for shape in shapes:
        current = path + [shape.name]
        yield shape, current
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes, current)


def finite_extrema(image: Image.Image) -> tuple[float, float]:
    extrema = image.getextrema()
    if isinstance(extrema[0], tuple):
        raise ValueError("Expected a scalar image")
    low, high = float(extrema[0]), float(extrema[1])
    if math.isfinite(low) and math.isfinite(high):
        return low, high
    values = (float(value) for value in image.getdata())
    finite_values = [value for value in values if math.isfinite(value)]
    if not finite_values:
        raise ValueError("No finite pixels")
    return min(finite_values), max(finite_values)


def scalar_preview(image: Image.Image) -> tuple[Image.Image, dict[str, Any]]:
    float_image = image.convert("F")
    low, high = finite_extrema(float_image)
    if high == low:
        preview = Image.new("L", image.size, 0)
    else:
        scale = 255.0 / (high - low)
        preview = float_image.point(lambda value: (value - low) * scale).convert("L")
    return preview, {
        "mapping": "linear-minmax",
        "display_min": low,
        "display_max": high,
        "quantitative_use": False,
    }


def color_preview(image: Image.Image) -> tuple[Image.Image, dict[str, Any]]:
    if image.mode == "P":
        return image.convert("RGBA"), {"mapping": "palette-to-rgba", "quantitative_use": False}
    if image.mode == "CMYK":
        icc = image.info.get("icc_profile")
        if icc:
            try:
                source_profile = ImageCms.ImageCmsProfile(io.BytesIO(icc))
                target_profile = ImageCms.createProfile("sRGB")
                converted = ImageCms.profileToProfile(image, source_profile, target_profile, outputMode="RGB")
                return converted, {"mapping": "icc-to-srgb", "quantitative_use": False}
            except Exception:
                pass
        return image.convert("RGB"), {"mapping": "cmyk-to-rgb", "quantitative_use": False}
    if image.mode in {"RGB", "RGBA", "L", "LA"}:
        return image.copy(), {"mapping": "identity", "quantitative_use": False}
    if image.mode in {"1", "I", "F"} or image.mode.startswith("I;16"):
        return scalar_preview(image)
    return image.convert("RGBA"), {"mapping": f"{image.mode}-to-rgba", "quantitative_use": False}


def tiff_record(blob: bytes, source_name: str, output_dir: Path) -> dict[str, Any]:
    digest = hashlib.sha256(blob).hexdigest()
    asset_dir = output_dir / digest
    asset_dir.mkdir(parents=True, exist_ok=True)
    original_path = asset_dir / "original.tif"
    if not original_path.exists():
        original_path.write_bytes(blob)

    frame_records = []
    with Image.open(io.BytesIO(blob)) as image:
        frame_count = int(getattr(image, "n_frames", 1))
        base_info = {
            "source_filename": source_name,
            "sha256": digest,
            "bytes": len(blob),
            "format": image.format,
            "size": list(image.size),
            "frame_count": frame_count,
            "compression": image.info.get("compression"),
            "dpi": [float(value) for value in image.info["dpi"]] if image.info.get("dpi") else None,
            "original": str(original_path.relative_to(output_dir)),
        }
        for frame_number in range(frame_count):
            image.seek(frame_number)
            frame = image.copy()
            preview, transform = color_preview(frame)
            preview_path = asset_dir / f"frame-{frame_number:03d}.png"
            preview.save(preview_path, format="PNG")
            frame_records.append(
                {
                    "frame": frame_number,
                    "source_mode": frame.mode,
                    "size": list(frame.size),
                    "preview": str(preview_path.relative_to(output_dir)),
                    "preview_mode": preview.mode,
                    "preview_transform": transform,
                }
            )
    base_info["frames"] = frame_records
    (asset_dir / "metadata.json").write_text(
        json.dumps(base_info, ensure_ascii=False, indent=2) + os.linesep,
        encoding="utf-8",
    )
    return base_info


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    if not args.pptx.is_file():
        parser.error(f"PPTX not found: {args.pptx}")
    args.output_dir.mkdir(parents=True, exist_ok=True)

    prs = Presentation(args.pptx)
    assets: dict[str, dict[str, Any]] = {}
    warnings: list[dict[str, Any]] = []
    for slide_number, slide in enumerate(prs.slides, 1):
        for shape, shape_path in walk(slide.shapes, []):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            try:
                part = image_part(shape)
                if part is None:
                    continue
                content_type = getattr(part, "content_type", "")
                partname = str(getattr(part, "partname", ""))
                if content_type != "image/tiff" and not partname.lower().endswith((".tif", ".tiff")):
                    continue
                blob = part.blob
                digest = hashlib.sha256(blob).hexdigest()
                if digest not in assets:
                    assets[digest] = tiff_record(blob, Path(partname).name, args.output_dir)
                    assets[digest]["occurrences"] = []
                assets[digest]["occurrences"].append(
                    {
                        "slide": slide_number,
                        "shape_path": shape_path,
                        "media_path": partname,
                    }
                )
            except Exception as exc:
                warnings.append(
                    {
                        "severity": "error",
                        "slide": slide_number,
                        "shape_path": shape_path,
                        "detail": f"{type(exc).__name__}: {exc}",
                    }
                )

    manifest = {
        "schema_version": 1,
        "source": {
            "path": str(args.pptx.resolve()),
            "size_bytes": args.pptx.stat().st_size,
            "modified_time_ns": args.pptx.stat().st_mtime_ns,
        },
        "asset_count": len(assets),
        "occurrence_count": sum(len(record["occurrences"]) for record in assets.values()),
        "assets": list(assets.values()),
        "warnings": warnings,
    }
    manifest_path = args.output_dir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + os.linesep, encoding="utf-8")
    print(json.dumps({"manifest": str(manifest_path), "assets": len(assets), "warnings": len(warnings)}))
    return 0 if not warnings else 1


if __name__ == "__main__":
    raise SystemExit(main())
