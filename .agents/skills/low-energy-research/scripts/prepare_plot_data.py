"""Read an XLSX long table; preserve all selected numeric observations as wide plot data."""
import argparse
import hashlib
import json
import math
from pathlib import Path
import openpyxl


def number(value, where, allow_blank=False):
    if value is None and allow_blank:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"Expected a finite number at {where}: {value!r}")
    return value


def prepare(source, config):
    source = Path(source).resolve()
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    wf = openpyxl.load_workbook(source, read_only=True, data_only=False)
    try:
        ws = wb[config['sheet']]
        headers = list(next(ws.values))
        fields = [config['series'], config['x']] + [m['column'] for m in config['metrics']]
        if not config['metrics']:
            raise ValueError('At least one metric is required')
        for field in fields:
            if headers.count(field) != 1:
                raise ValueError(f'Missing or ambiguous source column: {field}')
        indexes = {field: headers.index(field) for field in fields}
        names = [m['sheet'] for m in config['metrics']]
        if len(set(names)) != len(names) or any(not n or len(n) > 31 or any(c in n for c in '[]:*?/\\') for n in names):
            raise ValueError('Invalid or duplicate output sheet names')
        selected = config.get('series_values')
        groups = {}
        formulas = wf[config['sheet']].iter_rows(min_row=2, values_only=True)
        for rownum, (row, frow) in enumerate(zip(ws.iter_rows(min_row=2, values_only=True), formulas), 2):
            if all(v is None for v in row):
                continue
            key = row[indexes[config['series']]]
            if key is None or str(key).strip() == '':
                raise ValueError(f'Missing series identity on row {rownum}')
            key = str(key)
            if selected is not None and key not in selected:
                continue
            for field in fields:
                i = indexes[field]
                if row[i] is None and isinstance(frow[i], str) and frow[i].startswith('='):
                    raise ValueError(f'Uncached formula at row {rownum}, column {field}')
            x = number(row[indexes[config['x']]], f'row {rownum} X')
            ys = [number(row[indexes[m['column']]], f'row {rownum} {m["column"]}', True) for m in config['metrics']]
            groups.setdefault(key, []).append((x, ys, rownum))
        if not groups:
            raise ValueError('No selected observations')
        if selected is not None:
            if len(set(selected)) != len(selected) or set(selected) != set(groups):
                raise ValueError('Requested series are missing or duplicated')
            groups = {k: groups[k] for k in selected}
        for key, rows in groups.items():
            rows.sort(key=lambda r: r[0])
            if len({r[0] for r in rows}) != len(rows):
                raise ValueError(f'Duplicate X values within {key}; resolve series identity first')
        keys = list(groups)
        xs = [[r[0] for r in rows] for rows in groups.values()]
        shared = all(x == xs[0] for x in xs[1:])
        tables = []
        for j, metric in enumerate(config['metrics']):
            if shared:
                header = [config['x_label']] + [f'{k} | {metric["label"]}' for k in keys]
                values = [[x] + [groups[k][i][1][j] for k in keys] for i, x in enumerate(xs[0])]
            else:
                header = [v for k in keys for v in (f'{k} | {config["x_label"]}', f'{k} | {metric["label"]}')]
                values = []
                for i in range(max(map(len, xs))):
                    values.append([v for k in keys for v in ((groups[k][i][0], groups[k][i][1][j]) if i < len(groups[k]) else (None, None))])
            tables.append({'name': metric['sheet'], 'header': header, 'values': values})
        return {'source': str(source), 'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
                'recipe': config, 'shared_x': shared, 'series': keys,
                'source_rows': {k: [r[2] for r in rows] for k, rows in groups.items()}, 'tables': tables}
    finally:
        wb.close()
        wf.close()


if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--input', required=True)
    p.add_argument('--recipe', required=True)
    p.add_argument('--output', required=True)
    a = p.parse_args()
    result = prepare(a.input, json.loads(Path(a.recipe).read_text(encoding='utf-8-sig')))
    target = Path(a.output)
    if target.resolve() == Path(a.input).resolve():
        raise ValueError('Output must not overwrite input')
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(result, ensure_ascii=False, allow_nan=False), encoding='utf-8')
    print(json.dumps({'series': result['series'], 'shared_x': result['shared_x'], 'tables': [{'name': t['name'], 'rows': len(t['values'])} for t in result['tables']]}, ensure_ascii=False))
